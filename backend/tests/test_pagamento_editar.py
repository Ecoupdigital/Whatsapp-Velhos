"""Editar pagamento de evento: o valor novo chega em quem estava vinculado."""
from models import BaileVinculo, Evento, EventoParticipante, EventoPatrocinio, Transacao


def _evento_com_participante(TestingSession, valor=2340.0):
    db = TestingSession()
    ev = Evento(tipo="outro", titulo="Baile", status="em_andamento")
    db.add(ev)
    db.commit()
    p = EventoParticipante(evento_id=ev.id, nome_avulso="Ricardo", valor=valor, valor_pago=0)
    db.add(p)
    db.commit()
    ids = ev.id, p.id
    db.close()
    return ids


def _pagar(client, evento_id, participante_id, valor):
    r = client.post(
        f"/api/eventos/{evento_id}/participantes/{participante_id}/pagamento",
        json={"valor": valor, "data": "2026-09-26"},
    )
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _participante(client, evento_id, participante_id):
    r = client.get(f"/api/eventos/{evento_id}/participantes/{participante_id}")
    assert r.status_code == 200, r.text
    return r.json()


def test_editar_lancamento_no_financeiro_atualiza_o_evento(client, TestingSession):
    evento_id, part_id = _evento_com_participante(TestingSession)
    tx_id = _pagar(client, evento_id, part_id, 570)
    assert _participante(client, evento_id, part_id)["valor_pago"] == 570

    r = client.put(f"/api/transacoes/{tx_id}", json={"valor": 540})
    assert r.status_code == 200, r.text

    assert _participante(client, evento_id, part_id)["valor_pago"] == 540
    resumo = client.get(f"/api/eventos/{evento_id}/resumo").json()
    assert resumo["valor_arrecadado"] == 540


def test_editar_pagamento_pelo_evento(client, TestingSession):
    evento_id, part_id = _evento_com_participante(TestingSession)
    tx_id = _pagar(client, evento_id, part_id, 570)

    r = client.put(
        f"/api/eventos/{evento_id}/pagamentos/{tx_id}",
        json={"valor": 540, "data": "2026-09-20"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["valor"] == 540
    assert r.json()["data"] == "2026-09-20"

    p = _participante(client, evento_id, part_id)
    assert p["valor_pago"] == 540
    assert p["data_pagamento"] == "2026-09-20"

    db = TestingSession()
    assert db.query(Transacao).filter(Transacao.id == tx_id).one().valor == 540
    db.close()


def test_editar_pagamento_nao_passa_do_que_falta(client, TestingSession):
    evento_id, part_id = _evento_com_participante(TestingSession, valor=600)
    _pagar(client, evento_id, part_id, 100)
    tx_id = _pagar(client, evento_id, part_id, 400)

    r = client.put(f"/api/eventos/{evento_id}/pagamentos/{tx_id}", json={"valor": 501})
    assert r.status_code == 400
    assert "500.00" in r.json()["detail"]

    ok = client.put(f"/api/eventos/{evento_id}/pagamentos/{tx_id}", json={"valor": 500})
    assert ok.status_code == 200, ok.text
    p = _participante(client, evento_id, part_id)
    assert p["valor_pago"] == 600
    assert p["pago"] == 1


def test_editar_lancamento_de_patrocinio_ajusta_o_vinculo(client, TestingSession):
    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    pat = EventoPatrocinio(evento_id=ev.id, nome_patrocinador="Don Vicente", valor=60, pago="nao")
    db.add(pat)
    db.commit()
    evento_id, pat_id = ev.id, pat.id
    db.close()

    r = client.post(
        f"/api/eventos/{evento_id}/baile/patrocinios/{pat_id}/pagamento",
        json={"valor": 60, "data": "2026-09-25"},
    )
    assert r.status_code == 200, r.text
    db = TestingSession()
    tx_id = db.query(Transacao).filter(Transacao.categoria == "patrocinio").one().id
    db.close()

    r = client.put(f"/api/transacoes/{tx_id}", json={"valor": 50})
    assert r.status_code == 200, r.text

    db = TestingSession()
    assert db.query(BaileVinculo).filter(BaileVinculo.transacao_id == tx_id).one().valor == 50
    assert db.query(EventoPatrocinio).get(pat_id).pago == "a_confirmar"
    db.close()


def test_lancamento_dividido_nao_baixa_abaixo_das_partes(client, TestingSession):
    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    a = EventoPatrocinio(evento_id=ev.id, nome_patrocinador="A", valor=60, pago="nao")
    b = EventoPatrocinio(evento_id=ev.id, nome_patrocinador="B", valor=60, pago="nao")
    db.add_all([a, b])
    db.commit()
    tx = Transacao(tipo="entrada", categoria="patrocinio", descricao="PIX", valor=120, data="2026-09-25")
    db.add(tx)
    db.commit()
    evento_id, a_id, b_id, tx_id = ev.id, a.id, b.id, tx.id
    db.close()

    for pat_id in (a_id, b_id):
        r = client.post(
            f"/api/eventos/{evento_id}/baile/vincular",
            json={"transacao_id": tx_id, "patrocinio_id": pat_id},
        )
        assert r.status_code == 200, r.text

    r = client.put(f"/api/transacoes/{tx_id}", json={"valor": 100})
    assert r.status_code == 400
    assert "desvincule" in r.json()["detail"].lower()

    db = TestingSession()
    assert db.query(Transacao).get(tx_id).valor == 120
    db.close()


def test_importar_planilha_puxa_o_pago_do_lancamento(client, TestingSession, tmp_path):
    """Pago guardado zerado com lançamento ligado: o import acerta pelo lançamento."""
    from openpyxl import Workbook

    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    p = EventoParticipante(evento_id=ev.id, nome_avulso="Nando", valor=900, valor_pago=0)
    db.add(p)
    db.commit()
    db.add(Transacao(
        tipo="entrada", categoria="cartao_baile", descricao="Nando 5 cartoes",
        valor=900, data="2026-09-20", evento_id=ev.id, evento_participante_id=p.id,
    ))
    db.commit()
    evento_id, part_id = ev.id, p.id
    db.close()

    wb = Workbook()
    params = wb.active
    params.title = "Parâmetros"
    params["B6"] = 180
    params["B8"] = 30
    params["B17"] = 90
    cartoes = wb.create_sheet("Cartões")
    cartoes["B5"] = "Nando"
    cartoes["E5"] = 5
    wb.create_sheet("Patrocínios")
    caminho = tmp_path / "baile.xlsx"
    wb.save(caminho)
    with open(caminho, "rb") as f:
        r = client.post(
            f"/api/eventos/{evento_id}/baile/importar",
            files={"file": ("baile.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert r.status_code == 200, r.text

    p = _participante(client, evento_id, part_id)
    assert p["valor_pago"] == 900
    assert p["pago"] == 1
