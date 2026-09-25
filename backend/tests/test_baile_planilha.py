"""Planilha do baile: meio cartao, lucro a parte, e vinculo que nao cria lancamento."""
from pathlib import Path

from openpyxl import Workbook

from models import BaileVinculo, Evento, EventoParticipante, EventoPatrocinio, Jogador, Transacao


def _planilha(path: Path):
    wb = Workbook()
    params = wb.active
    params.title = "Parâmetros"
    params["B6"] = 180
    params["B8"] = 30
    params["B12"] = 17100
    params["B17"] = 90
    cartoes = wb.create_sheet("Cartões")
    cartoes["B5"] = "Tucum"
    cartoes["C5"] = "51 9810-1638"
    cartoes["D5"] = "Respondeu"
    cartoes["E5"] = 3
    cartoes["I5"] = 2
    cartoes["K5"] = "106 a 110"
    cartoes["L5"] = "Não"
    cartoes["M5"] = "Lucro de 2 cartões"
    cartoes["B6"] = "Amauri"
    cartoes["C6"] = "51 9989-2263"
    cartoes["D6"] = "Respondeu"
    cartoes["E6"] = 7.5
    cartoes["I6"] = 0
    cartoes["K6"] = "111 a 115"
    cartoes["L6"] = "A confirmar"
    pats = wb.create_sheet("Patrocínios")
    pats["B4"] = "Amauri"
    pats["C4"] = "Bruxel"
    pats["D4"] = 60
    pats["E4"] = "Sim"
    pats["F4"] = "Sim"
    wb.save(path)


def test_importar_nao_cria_lancamento_e_aceita_meio_cartao(client, TestingSession, tmp_path):
    db = TestingSession()
    db.add(Jogador(nome="Tucum", apelido="Tucum", telefone="555198101638", tipo="jogador"))
    db.add(Jogador(nome="Amauri", apelido="Amauri", telefone="555199892263", tipo="socio"))
    ev = Evento(tipo="baile", titulo="Baile", status="planejado")
    db.add(ev)
    db.commit()
    evento_id = ev.id
    antes = db.query(Transacao).count()
    db.close()

    caminho = tmp_path / "baile.xlsx"
    _planilha(caminho)
    with open(caminho, "rb") as f:
        r = client.post(
            f"/api/eventos/{evento_id}/baile/importar",
            files={"file": ("baile.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert r.status_code == 200, r.text
    assert r.json()["atletas"] == 2
    assert r.json()["patrocinios"] == 1

    visao = client.get(f"/api/eventos/{evento_id}/baile")
    assert visao.status_code == 200
    body = visao.json()
    por_nome = {p["nome"]: p for p in body["participantes"]}
    assert por_nome["Tucum"]["qtd_venda"] == 3
    assert por_nome["Tucum"]["qtd_lucro"] == 2
    assert por_nome["Tucum"]["receita_cartao"] == 540
    assert por_nome["Tucum"]["receita_lucro"] == 180
    assert por_nome["Tucum"]["valor"] == 720
    assert por_nome["Tucum"]["valor_pago"] == 0
    assert por_nome["Tucum"]["numero_inicio"] == 106
    assert por_nome["Amauri"]["qtd_venda"] == 7.5
    assert por_nome["Amauri"]["acerto"] == "a_confirmar"
    assert body["patrocinios"][0]["nome_patrocinador"] == "Bruxel"
    assert body["patrocinios"][0]["pago"] == "sim"
    assert body["patrocinios"][0]["valor_vinculado"] == 0

    db = TestingSession()
    assert db.query(Transacao).count() == antes
    db.close()


def test_vincular_nao_duplica_o_lancamento(client, TestingSession):
    db = TestingSession()
    db.add(Jogador(nome="Mauro", apelido="Mauro", telefone="555199211910", tipo="jogador"))
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    j = db.query(Jogador).filter(Jogador.nome == "Mauro").one()
    p = EventoParticipante(
        evento_id=ev.id, jogador_id=j.id, qtd_venda=5, qtd_lucro=0,
        status="confirmado", valor=900, valor_pago=0, pago=0,
    )
    db.add(p)
    db.flush()
    tx = Transacao(
        tipo="entrada", categoria="cartao_baile", descricao="Mauro 5 Cartoes Baile",
        valor=900, data="2026-09-19",
    )
    db.add(tx)
    db.commit()
    evento_id, pid, txid = ev.id, p.id, tx.id
    db.close()

    r = client.post(
        f"/api/eventos/{evento_id}/baile/vincular",
        json={"transacao_id": txid, "participante_id": pid},
    )
    assert r.status_code == 200, r.text
    linha = next(x for x in r.json()["participantes"] if x["id"] == pid)
    assert linha["valor_pago"] == 900
    assert linha["valor"] == 900

    db = TestingSession()
    assert db.query(Transacao).count() == 1
    gravado = db.query(Transacao).one()
    assert gravado.valor == 900
    vinculo = db.query(BaileVinculo).one()
    assert vinculo.participante_id == pid
    assert vinculo.valor == 900
    db.close()

    de_novo = client.post(
        f"/api/eventos/{evento_id}/baile/vincular",
        json={"transacao_id": txid, "participante_id": pid},
    )
    assert de_novo.status_code == 200
    db = TestingSession()
    assert db.query(Transacao).count() == 1
    db.close()


def test_cria_e_marca_logo_do_patrocinio(client, TestingSession):
    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    evento_id = ev.id
    db.close()

    criado = client.post(
        f"/api/eventos/{evento_id}/baile/patrocinios",
        json={"nome_jogador": "Jonathan", "nome_patrocinador": "Don Vicente", "valor": 60},
    )
    assert criado.status_code == 201, criado.text
    pat = criado.json()["patrocinios"][0]
    assert pat["pago"] == "nao"
    assert pat["logo_enviado"] == "nao"

    atualizado = client.put(
        f"/api/eventos/{evento_id}/baile/patrocinios/{pat['id']}",
        json={"logo_enviado": "sim", "pago": "sim"},
    )
    assert atualizado.status_code == 200, atualizado.text
    pat = atualizado.json()["patrocinios"][0]
    assert pat["logo_enviado"] == "sim"
    assert pat["pago"] == "sim"
    assert atualizado.json()["resumo"]["patrocinios_marcados_pagos"] == 1


def test_um_pix_de_120_cobre_dois_patrocinios(client, TestingSession):
    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    a = EventoPatrocinio(evento_id=ev.id, nome_jogador="Jonathan", nome_patrocinador="Don Vicente", valor=60, pago="sim")
    b = EventoPatrocinio(evento_id=ev.id, nome_jogador="Jonathan", nome_patrocinador="Ecoinset", valor=60, pago="sim")
    db.add_all([a, b])
    tx = Transacao(
        tipo="entrada", categoria="patrocinio", descricao="2 Patrocinios Baile",
        valor=120, data="2026-09-17",
    )
    db.add(tx)
    db.commit()
    evento_id, a_id, b_id, tx_id = ev.id, a.id, b.id, tx.id
    db.close()

    r1 = client.post(
        f"/api/eventos/{evento_id}/baile/vincular",
        json={"transacao_id": tx_id, "patrocinio_id": a_id},
    )
    assert r1.status_code == 200, r1.text
    body = r1.json()
    por_id = {p["id"]: p for p in body["patrocinios"]}
    assert por_id[a_id]["valor_vinculado"] == 60
    assert por_id[b_id]["valor_vinculado"] == 0
    lanc = body["lancamentos"][0]
    assert lanc["restante"] == 60

    r2 = client.post(
        f"/api/eventos/{evento_id}/baile/vincular",
        json={"transacao_id": tx_id, "patrocinio_id": b_id},
    )
    assert r2.status_code == 200, r2.text
    body = r2.json()
    por_id = {p["id"]: p for p in body["patrocinios"]}
    assert por_id[a_id]["valor_vinculado"] == 60
    assert por_id[b_id]["valor_vinculado"] == 60
    assert body["resumo"]["caixa_patrocinio"] == 120
    assert body["lancamentos"][0]["restante"] == 0

    db = TestingSession()
    gravado = db.query(Transacao).one()
    assert gravado.valor == 120
    assert db.query(Transacao).count() == 1
    db.close()


def test_vinculo_antigo_de_120_deixa_o_segundo_livre(client, TestingSession):
    db = TestingSession()
    ev = Evento(tipo="baile", titulo="Baile", status="em_andamento", valor_cartao=180, valor_lucro=90)
    db.add(ev)
    db.commit()
    a = EventoPatrocinio(evento_id=ev.id, nome_jogador="Caue", nome_patrocinador="Glam", valor=60, pago="sim")
    b = EventoPatrocinio(evento_id=ev.id, nome_jogador="Caue", nome_patrocinador="Pioner", valor=60, pago="sim")
    db.add_all([a, b])
    db.flush()
    tx = Transacao(
        tipo="entrada", categoria="patrocinio", descricao="Caue 2 patrocinios",
        valor=120, data="2026-09-19", evento_id=ev.id, patrocinio_id=a.id,
    )
    db.add(tx)
    db.commit()
    evento_id, a_id, b_id, tx_id = ev.id, a.id, b.id, tx.id
    db.close()

    visao = client.get(f"/api/eventos/{evento_id}/baile").json()
    por_id = {p["id"]: p for p in visao["patrocinios"]}
    assert por_id[a_id]["valor_vinculado"] == 60
    assert visao["lancamentos"][0]["restante"] == 60

    r = client.post(
        f"/api/eventos/{evento_id}/baile/vincular",
        json={"transacao_id": tx_id, "patrocinio_id": b_id},
    )
    assert r.status_code == 200, r.text
    por_id = {p["id"]: p for p in r.json()["patrocinios"]}
    assert por_id[a_id]["valor_vinculado"] == 60
    assert por_id[b_id]["valor_vinculado"] == 60
    assert r.json()["resumo"]["caixa_patrocinio"] == 120
