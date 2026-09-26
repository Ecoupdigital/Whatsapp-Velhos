"""Importa a planilha do baile para o evento e recalcula o que cada um deve.

Nao cria lancamento financeiro. Quem ja pagou continua no caixa e so e
vinculado depois, um lancamento por vez.
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models import Evento, EventoCartaoFaixa, EventoParticipante, EventoPatrocinio, Jogador
from routers.eventos import _recalc_recebidos, _recalcular_valor_esperado, _recalcular_valor_pago

PLANILHA_PADRAO = Path(
    "/home/vault/05-pessoal/Velhos parceiros/Controle Baile 2026 - Cartoes e Patrocinio - atualizada.xlsx"
)

# Nome curto da aba de patrocinio -> nome do atleta na aba de cartoes.
ALIASES = {
    "fiuza": "devair",
    "fabio": "fabio krein",
    "fábio": "fabio krein",
    "seco": "marcelo seco",
    "marquinhos": "marcos krein",
    "caue": "cauê",
    "cauê": "cauê",
    "nando": "luis fernando",
    "cris": "cristian rosa",
    "jonathan renan": "jonathan renan",
    "jonathan k": "jonathan kolling",
    "jonathan kolling": "jonathan kolling",
}


def _sem_acento(s: str) -> str:
    nfkd = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in nfkd if not unicodedata.combining(c)).lower().strip()


def _chave_nome(s: str) -> str:
    base = _sem_acento(s or "")
    base = re.sub(r"\(.*?\)", " ", base)
    base = re.sub(r"[^a-z0-9 ]", " ", base)
    base = re.sub(r"\s+", " ", base).strip()
    return ALIASES.get(base, base)


def _telefone(s: str | None) -> str:
    d = "".join(c for c in (s or "") if c.isdigit())
    if d.startswith("55") and len(d) > 11:
        d = d[2:]
    return d[-9:] if len(d) >= 9 else d


def _acerto(valor) -> str:
    t = _sem_acento(str(valor or ""))
    if t == "sim":
        return "sim"
    if "confirm" in t:
        return "a_confirmar"
    return "nao"


def _resposta(valor) -> str:
    t = _sem_acento(str(valor or ""))
    if t == "respondeu":
        return "respondeu"
    if "confirm" in t:
        return "a_confirmar"
    return "sem_resposta"


def _logo(valor) -> str:
    t = _sem_acento(str(valor or ""))
    if t == "sim":
        return "sim"
    if t == "pendente":
        return "pendente"
    return "nao"


def _faixa(texto) -> tuple[int | None, int | None]:
    if not texto:
        return None, None
    m = re.search(r"(\d+)\s*a\s*(\d+)", str(texto))
    if not m:
        return None, None
    return int(m.group(1)), int(m.group(2))


def _celula(v):
    if v is None:
        return None
    if hasattr(v, "text"):
        return v.text
    return v


def _indexar_jogadores(db: Session) -> tuple[dict, dict]:
    por_tel = {}
    por_nome = {}
    for j in db.query(Jogador).all():
        tel = _telefone(j.telefone)
        if tel:
            por_tel.setdefault(tel, j)
        for pedaco in (j.nome, j.apelido):
            if pedaco:
                por_nome.setdefault(_chave_nome(pedaco), j)
    return por_tel, por_nome


def _achar_jogador(por_tel, por_nome, nome, telefone):
    tel = _telefone(telefone)
    if tel and tel in por_tel:
        return por_tel[tel]
    chave = _chave_nome(nome)
    if chave in por_nome:
        return por_nome[chave]
    # primeiro nome, se for unico o bastante
    primeiro = chave.split(" ")[0] if chave else ""
    if primeiro and primeiro in por_nome:
        return por_nome[primeiro]
    return None


def _garantir_faixa(db: Session, p: EventoParticipante, inicio, fim):
    if inicio is None or fim is None:
        return
    faixa = next((f for f in p.faixas if not f.sem_numero), None)
    if faixa is None:
        faixa = EventoCartaoFaixa(evento_participante_id=p.id, sem_numero=0)
        db.add(faixa)
        p.faixas.append(faixa)
    faixa.numero_inicio = inicio
    faixa.numero_fim = fim
    faixa.quantidade = fim - inicio + 1
    faixa.sem_numero = 0
    p.numero_inicio = inicio
    p.numero_fim = fim
    _recalc_recebidos(p)


def importar_arquivo(db: Session, evento_id: int, caminho: str | Path) -> dict:
    evento = db.query(Evento).filter(Evento.id == evento_id).first()
    if not evento:
        raise ValueError("Evento nao encontrado")

    wb = load_workbook(caminho, data_only=True)
    cartoes = wb["Cartões"]
    patrocinios = wb["Patrocínios"]
    params = wb["Parâmetros"]
    valor_cartao = float(_celula(params["B6"].value) or 180)
    valor_lucro = float(_celula(params["B17"].value) or 90)
    valor_brinde = float(_celula(params["B8"].value) or 30)

    evento.valor_cartao = valor_cartao
    evento.valor_lucro = valor_lucro
    evento.valor_brinde = valor_brinde
    if not evento.custo_estimado:
        evento.custo_estimado = float(_celula(params["B12"].value) or 0)
    if evento.status == "planejado":
        evento.status = "em_andamento"

    por_tel, por_nome = _indexar_jogadores(db)
    atletas = {}
    sem_cadastro = []
    atualizados = 0

    for r in range(5, cartoes.max_row + 1):
        nome = _celula(cartoes.cell(r, 2).value)
        if not nome or str(nome).strip().upper() == "TOTAL":
            continue
        nome = str(nome).strip()
        tel = _celula(cartoes.cell(r, 3).value)
        jogador = _achar_jogador(por_tel, por_nome, nome, tel)
        if jogador is None and nome not in ("Vinícius B. Boesing",):
            # Vinicius pode nao estar no elenco. Os demais tambem seguem como avulso se nao casar.
            sem_cadastro.append(nome)

        q = db.query(EventoParticipante).filter(EventoParticipante.evento_id == evento_id)
        if jogador is not None:
            p = q.filter(EventoParticipante.jogador_id == jogador.id).first()
        else:
            p = q.filter(EventoParticipante.nome_avulso == nome).first()
        if p is None:
            p = EventoParticipante(
                evento_id=evento_id,
                jogador_id=jogador.id if jogador else None,
                nome_avulso=None if jogador else nome,
                status="confirmado",
                pago=0,
                valor=0,
                valor_pago=0,
            )
            db.add(p)
            db.flush()
        p.qtd_venda = float(_celula(cartoes.cell(r, 5).value) or 0)
        p.qtd_lucro = float(_celula(cartoes.cell(r, 9).value) or 0)
        p.status_resposta = _resposta(_celula(cartoes.cell(r, 4).value))
        p.acerto = _acerto(_celula(cartoes.cell(r, 12).value))
        obs = _celula(cartoes.cell(r, 13).value)
        if obs:
            p.observacoes = str(obs)
        inicio, fim = _faixa(_celula(cartoes.cell(r, 11).value))
        db.flush()
        db.refresh(p)
        _garantir_faixa(db, p, inicio, fim)
        p.evento = evento
        _recalcular_valor_esperado(p, evento)
        _recalcular_valor_pago(db, p)
        atletas[_chave_nome(nome)] = p
        # tambem indexa o nome do cadastro, pra patrocinio achar "Seco" -> Marcelo Seco
        if jogador is not None:
            atletas.setdefault(_chave_nome(jogador.nome or ""), p)
            if jogador.apelido:
                atletas.setdefault(_chave_nome(jogador.apelido), p)
        atualizados += 1

    pats = 0
    for r in range(4, patrocinios.max_row + 1):
        quem = _celula(patrocinios.cell(r, 2).value)
        marca = _celula(patrocinios.cell(r, 3).value)
        if not quem or not marca or str(quem).strip().upper() == "TOTAL":
            continue
        quem = str(quem).strip()
        marca = str(marca).strip()
        chave = _chave_nome(quem)
        p = atletas.get(chave)
        existente = (
            db.query(EventoPatrocinio)
            .filter(
                EventoPatrocinio.evento_id == evento_id,
                EventoPatrocinio.nome_jogador == quem,
                EventoPatrocinio.nome_patrocinador == marca,
            )
            .first()
        )
        data_pg = _celula(patrocinios.cell(r, 7).value)
        if hasattr(data_pg, "strftime"):
            data_pg = data_pg.strftime("%Y-%m-%d")
        elif data_pg is not None:
            data_pg = str(data_pg)
        campos = dict(
            participante_id=p.id if p else None,
            jogador_id=p.jogador_id if p else None,
            valor=float(_celula(patrocinios.cell(r, 4).value) or 60),
            logo_enviado=_logo(_celula(patrocinios.cell(r, 5).value)),
            pago=_acerto(_celula(patrocinios.cell(r, 6).value)),
            data_pagamento=data_pg,
            contato=_celula(patrocinios.cell(r, 8).value),
            observacoes=_celula(patrocinios.cell(r, 9).value),
        )
        if existente is None:
            db.add(EventoPatrocinio(
                evento_id=evento_id,
                nome_jogador=quem,
                nome_patrocinador=marca,
                **campos,
            ))
        else:
            for k, v in campos.items():
                setattr(existente, k, v)
        pats += 1

    db.commit()
    return {
        "atletas": atualizados,
        "patrocinios": pats,
        "sem_cadastro": sem_cadastro,
        "valor_cartao": valor_cartao,
        "valor_lucro": valor_lucro,
    }
